import os
import msal
import requests
import pandas as pd
import psycopg2
import io
import openpyxl
import time
from supabase import create_client
from concurrent.futures import ThreadPoolExecutor

# ── Credenciales ──────────────────────────────────────────
CLIENT_ID      = os.environ["MS_CLIENT_ID"]
TENANT_ID      = os.environ["MS_TENANT_ID"]
REFRESH_TOKEN  = os.environ["MS_REFRESH_TOKEN"]
PG_HOST        = os.environ["PG_HOST"]
PG_USER        = os.environ["PG_USER"]
PG_PASS        = os.environ["PG_PASS"]
PG_DB          = os.environ["PG_DB"]
SUPABASE_URL   = os.environ["SUPABASE_URL2"]
SUPABASE_KEY   = os.environ["SUPABASE_KEY2"]

SHAREPOINT_SITE = "escuelarefrigeracion.sharepoint.com"
SITE_PATH       = "/sites/ASESORASCOMERCIALES"
SUBFOLDER       = "2. BASE PROSPECTOS/BASE GENERAL"
SCOPES          = ["Sites.Read.All", "Files.Read.All"]

# IMPORTANTE: "Codigo" = código de PUBLICACIONES (NO se toca).
# "Codigo_Vendedor" = código de la asesora.
COLUMNAS = [
    "Ejecutivo", "Telefono", "Fechacreada", "Sede", "Programa",
    "Turno", "Codigo", "Canal", "Intervalo", "Medio",
    "Contacto", "Interesado", "Estado", "Objecion", "Observacion",
    "Codigo_Vendedor",
]

TABLAS = {
    "Base Carmen Montoya.xlsx":   "Base_Carmen",
    "Base Milagros Vargas.xlsx":  "Base_Gerson",
    "Base Diana Chavez.xlsx":     "Base_Diana",
    "Base Veronica La Rosa.xlsx": "Base_Veronica",
    "Base Dayana Balabarca.xlsx": "Base_Alonso22",
    "Base Sede Arequipa.xlsx":    "Base_Arequipa",     # nuevo
    "Base Sede Piura.xlsx":       "Base_Piura2",       # nuevo
    "Base Sede Trujillo.xlsx":    "Base_Trujillo2",    # nuevo
}

# ── CÓDIGO DE VENDEDOR POR TABLA (Excel) ───────────────────
# Para lo que viene de Excel el código se asigna por el ARCHIVO/TABLA (fijo),
# así un error de tipeo en el nombre del Ejecutivo NO afecta el código.
CODIGOS_TABLA = {
    "Base_Carmen":    "20",
    "Base_Gerson":    "16",
    "Base_Diana":     "19",
    "Base_Veronica":  "18",
    "Base_Alonso22":  "17",
    "Base_Arequipa":  "21",
    "Base_Piura2":    "23",
    "Base_Trujillo2": "22",
}

# ── CÓDIGO DE VENDEDOR POR NOMBRE (solo para Postgres/Chatwoot) ─────
# En Chatwoot no hay archivo del cual sacar el código → se usa el nombre del agente.
CODIGOS_NOMBRE = {
    "milagros vargas": "16",
    "balabarca, dayana yesenia": "17",
    "dayana balabarca": "17",
    "la rosa cardenas, veronica astrid": "18",
    "verónica la rosa": "18",
    "veronica la rosa": "18",
    "chavez padilla, diana": "19",
    "diana chávez": "19",
    "diana chavez": "19",
    "montoya dulanto, carmen isabel": "20",
    "carmen montoya": "20",
    "almendra peralta": "21",
    "peralta condori, almendra": "21",
    "diego lázaro": "22",
    "diego lazaro": "22",
    "lazaro quipuzco diego": "22",
    "juan carlos aguilar": "23",
    "aguilar ubillús, juan carlos": "23",
    "aguilar ubillus, juan carlos": "23",
}


def codigo_vendedor_por_nombre(nombre):
    """Código de vendedor según el nombre (para Postgres). '0' = no reconocido."""
    if not isinstance(nombre, str):
        return "0"
    return CODIGOS_NOMBRE.get(nombre.strip().lower(), "0")


# ── SharePoint ────────────────────────────────────────────
def get_access_token():
    app = msal.PublicClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}"
    )
    result = app.acquire_token_by_refresh_token(REFRESH_TOKEN, scopes=SCOPES)
    if "access_token" not in result:
        raise Exception(f"Error token: {result.get('error_description')}")
    return result["access_token"]


def get_site_id(token):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE}:{SITE_PATH}"
    r = requests.get(url, headers=headers)
    r.raise_for_status()
    return r.json()["id"]


def get_drive_id(token, site_id):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives"
    r = requests.get(url, headers=headers)
    r.raise_for_status()
    drives = r.json().get("value", [])
    for drive in drives:
        name_lower = drive["name"].lower()
        if "document" in name_lower or "compartid" in name_lower:
            return drive["id"]
    return drives[0]["id"]


def list_excel_files(token, drive_id):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root:/{SUBFOLDER}:/children"
    r = requests.get(url, headers=headers)
    r.raise_for_status()
    items = r.json().get("value", [])
    return [f for f in items if f["name"].endswith((".xlsx", ".xls"))]


def download_and_process(args):
    token, drive_id, file_id, file_name = args
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{file_id}/content"
    r = requests.get(url, headers=headers)
    r.raise_for_status()

    nombre_tabla = TABLAS.get(file_name)
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    df = None
    if nombre_tabla:
        for sheet in wb.worksheets:
            if nombre_tabla in sheet.tables:
                tabla = sheet.tables[nombre_tabla]
                data = list(sheet[tabla.ref])
                headers_row = [cell.value for cell in data[0]]
                rows = [[cell.value for cell in row] for row in data[1:]]
                df = pd.DataFrame(rows, columns=headers_row)
                break
    if df is None:
        df = pd.read_excel(io.BytesIO(r.content))

    for col in COLUMNAS:
        if col not in df.columns:
            df[col] = "-"

    if "Fechacreada" in df.columns:
        df["Fechacreada"] = pd.to_datetime(df["Fechacreada"], errors="coerce").dt.strftime("%d/%m/%Y")
        df["Fechacreada"] = df["Fechacreada"].fillna("-")

    # CÓDIGO POR TABLA: todas las filas de este Excel llevan el código del archivo.
    # (Robusto a errores de tipeo en la columna Ejecutivo.) NO se toca "Codigo".
    df["Codigo_Vendedor"] = CODIGOS_TABLA.get(nombre_tabla, "0")

    asignados = (df["Codigo_Vendedor"] != "0").sum()
    print(f"  ✅ {file_name}: {len(df)} filas ({asignados} con vendedor, {len(df)-asignados} en 0)")
    return df[COLUMNAS]


# ── PostgreSQL ────────────────────────────────────────────
def get_postgres_data():
    conn = psycopg2.connect(
        host=PG_HOST, user=PG_USER,
        password=PG_PASS, dbname=PG_DB, port=5432
    )
    query = """
        SELECT c.phone_number, c.created_at, u.name
        FROM contacts AS c
        LEFT JOIN conversations AS con ON c.id = con.contact_id
        LEFT JOIN users AS u ON con.assignee_id = u.id
        WHERE u.name IS NOT NULL
        ORDER BY c.created_at DESC
    """
    df = pd.read_sql(query, conn)
    conn.close()

    df["phone_number"] = df["phone_number"].astype(str)
    df["phone_number"] = df["phone_number"].str.replace(r"^\+51", "", regex=True)
    df["phone_number"] = df["phone_number"].str.replace("+", "", regex=False)
    df["created_at"] = pd.to_datetime(df["created_at"]).dt.strftime("%d/%m/%Y")

    df_mapped = pd.DataFrame("-", index=df.index, columns=COLUMNAS)
    df_mapped["Ejecutivo"]       = df["name"]
    df_mapped["Telefono"]        = df["phone_number"]
    df_mapped["Fechacreada"]     = df["created_at"]
    df_mapped["Canal"]           = "COPITO"
    # Postgres: código POR NOMBRE del agente (aquí no hay archivo/tabla).
    df_mapped["Codigo_Vendedor"] = df["name"].apply(codigo_vendedor_por_nombre)

    print(f"  ✅ PostgreSQL: {len(df_mapped)} filas")
    return df_mapped


# ── MAIN ──────────────────────────────────────────────────
inicio = time.time()
print("="*50)
print("🚀 Iniciando sincronización...")
print("="*50)

# 1. Preparar
print("\n🔑 Autenticando...")
token    = get_access_token()
site_id  = get_site_id(token)
drive_id = get_drive_id(token, site_id)
excels   = list_excel_files(token, drive_id)

# 2. Descargar Excel EN PARALELO
print("\n📁 Descargando Excel en paralelo...")
tareas = [
    (token, drive_id, f["id"], f["name"])
    for f in excels if f["name"] in TABLAS
]
all_dfs = []
with ThreadPoolExecutor(max_workers=5) as executor:
    resultados = list(executor.map(download_and_process, tareas))
    all_dfs.extend(resultados)

# 3. PostgreSQL
print("\n🗄️  Leyendo PostgreSQL...")
df_pg = get_postgres_data()
all_dfs.append(df_pg)

# 4. Unir todo
df_final = pd.concat(all_dfs, ignore_index=True)
df_final = df_final.fillna("-").astype(str).replace("nan", "-")
df_final["Codigo_Vendedor"] = df_final["Codigo_Vendedor"].replace(["-", "nan", ""], "0")

print(f"\n✅ Total unificado: {len(df_final)} filas")

# 5. Subir a Supabase con TRUNCATE
print("\n🗑️  Truncando tabla anterior...")
supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
supabase.rpc("truncate_datos_unificados").execute()

print("\n⬆️  Subiendo a Supabase...")
records = df_final.to_dict(orient="records")
batch_size = 1000
for i in range(0, len(records), batch_size):
    batch = records[i:i+batch_size]
    supabase.table("datos_unificados").insert(batch).execute()
    print(f"  📤 {min(i+batch_size, len(records))}/{len(records)}")

duracion = time.time() - inicio
print(f"\n🎉 Completado en {duracion:.1f} segundos")
